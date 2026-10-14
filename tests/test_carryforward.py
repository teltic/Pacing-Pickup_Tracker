import os
import tempfile
import unittest
from datetime import date

import openpyxl

from pacing_tracker.carryforward import (
    copy_pasted_sheets,
    find_latest_file,
    find_previous_file,
    header_column_map,
    load_previous_state,
    load_previous_state_for_folder,
    parse_filename_date,
)


def _write_fake_workbook(path, rows, signal_col=12, override_request_col=13, notes_col=14):
    """rows: list of (date, signal, override_request, notes). Column
    positions are adjustable (default: today's real layout, A/M/N/O) so
    tests can simulate a previous file saved under an OLDER column
    layout -- load_previous_state is supposed to find the right cells by
    header text regardless of position (2026-10-14 incident).
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Daily Pacing"
    headers = ["Date"] + [f"col{i}" for i in range(1, 22)]  # A..V header
    headers[signal_col] = "Signal"
    headers[override_request_col] = "Override Request"
    headers[notes_col] = "Notes"
    ws.append(headers)
    for d, signal, override_request, notes in rows:
        row = [d] + [None] * 21
        row[signal_col] = signal
        row[override_request_col] = override_request
        row[notes_col] = notes
        ws.append(row)
    wb.save(path)


class ParseFilenameDateTest(unittest.TestCase):
    def test_parses_reference_naming_convention(self):
        self.assertEqual(parse_filename_date("Daily_Pacing_Pickup_9.11.26.xlsx"), date(2026, 9, 11))

    def test_returns_none_for_non_matching_name(self):
        self.assertIsNone(parse_filename_date("something_else.xlsx"))


class FindPreviousFileTest(unittest.TestCase):
    def test_picks_most_recent_file_strictly_before_given_date(self):
        with tempfile.TemporaryDirectory() as tmp:
            for name in ["Daily_Pacing_Pickup_9.10.26.xlsx", "Daily_Pacing_Pickup_9.11.26.xlsx", "Daily_Pacing_Pickup_9.12.26.xlsx"]:
                open(os.path.join(tmp, name), "w").close()
            found = find_previous_file(tmp, date(2026, 9, 12))
            self.assertEqual(os.path.basename(found), "Daily_Pacing_Pickup_9.11.26.xlsx")

    def test_returns_none_when_no_earlier_file_exists(self):
        with tempfile.TemporaryDirectory() as tmp:
            open(os.path.join(tmp, "Daily_Pacing_Pickup_9.12.26.xlsx"), "w").close()
            self.assertIsNone(find_previous_file(tmp, date(2026, 9, 12)))


class FindLatestFileTest(unittest.TestCase):
    def test_picks_the_most_recent_by_filename_date(self):
        with tempfile.TemporaryDirectory() as tmp:
            for name in ["Daily_Pacing_Pickup_9.10.26.xlsx", "Daily_Pacing_Pickup_9.12.26.xlsx", "Daily_Pacing_Pickup_9.11.26.xlsx"]:
                open(os.path.join(tmp, name), "w").close()
            found = find_latest_file(tmp)
            self.assertEqual(os.path.basename(found), "Daily_Pacing_Pickup_9.12.26.xlsx")

    def test_returns_none_when_folder_is_empty(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertIsNone(find_latest_file(tmp))

    def test_ignores_files_that_dont_match_the_naming_convention(self):
        with tempfile.TemporaryDirectory() as tmp:
            open(os.path.join(tmp, "Daily_Pacing_Pickup_9.12.26.xlsx"), "w").close()
            open(os.path.join(tmp, "some_other_file.xlsx"), "w").close()
            found = find_latest_file(tmp)
            self.assertEqual(os.path.basename(found), "Daily_Pacing_Pickup_9.12.26.xlsx")


class LoadPreviousStateTest(unittest.TestCase):
    def test_reads_signal_override_and_notes_by_date(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "prev.xlsx")
            _write_fake_workbook(
                path,
                [
                    (date(2026, 9, 11), "▼ Behind ⚡ Spike", -0.1, "9/11 - behind"),
                    (date(2026, 9, 12), "✓ Booked", None, None),
                ],
            )
            state = load_previous_state(path)
            self.assertEqual(state["2026-09-11"]["signal"], "▼ Behind ⚡ Spike")
            self.assertEqual(state["2026-09-11"]["override_request"], -0.1)
            self.assertEqual(state["2026-09-11"]["notes"], "9/11 - behind")
            self.assertEqual(state["2026-09-12"]["signal"], "✓ Booked")
            self.assertIsNone(state["2026-09-12"]["override_request"])

    def test_for_folder_returns_empty_dict_when_no_history(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(load_previous_state_for_folder(tmp, date(2026, 9, 12)), {})

    def test_finds_columns_by_header_even_when_a_previous_file_used_an_older_layout(self):
        # 2026-10-14 incident: a column reorder shipped, and the very next
        # day's run read yesterday's (old-layout) file with today's fixed
        # indices -- Signal's new position held yesterday's Override
        # Request (a float), and _new_since_last_review_formula crashed
        # trying to .replace() it as text. This is the regression test:
        # Signal/Override Request/Notes sit at DIFFERENT columns than
        # today's real layout, and load_previous_state must still find
        # them correctly by header text, not position.
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "prev.xlsx")
            _write_fake_workbook(
                path,
                [(date(2026, 9, 11), "▼ Behind ⚡ Spike", -0.1, "9/11 - behind")],
                signal_col=11, override_request_col=12, notes_col=13,  # yesterday's (old) layout
            )
            state = load_previous_state(path)
            self.assertEqual(state["2026-09-11"]["signal"], "▼ Behind ⚡ Spike")
            self.assertEqual(state["2026-09-11"]["override_request"], -0.1)
            self.assertEqual(state["2026-09-11"]["notes"], "9/11 - behind")

    def test_raises_a_clear_error_when_a_required_header_is_missing(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "prev.xlsx")
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Daily Pacing"
            ws.append(["Date", "Weekday"])  # no Signal/Override Request/Notes at all
            wb.save(path)
            with self.assertRaises(ValueError) as ctx:
                load_previous_state(path)
            self.assertIn("Signal", str(ctx.exception))


class HeaderColumnMapTest(unittest.TestCase):
    def test_maps_header_text_to_zero_based_index(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Date", "Weekday", "Signal"])
        self.assertEqual(header_column_map(ws), {"Date": 0, "Weekday": 1, "Signal": 2})


class CopyPastedSheetsTest(unittest.TestCase):
    def test_copies_every_cell_from_the_previous_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            prev_path = os.path.join(tmp, "prev.xlsx")
            prev_wb = openpyxl.Workbook()
            prev_ws = prev_wb.active
            prev_ws.title = "Neighborhood Data - Mkt Occ"
            prev_ws["A1"] = "Date"
            prev_ws["B1"] = "Market Occupancy"
            prev_ws["A2"] = "2026-09-12"
            prev_ws["B2"] = 42.0
            prev_wb.create_sheet("Neighborhood Data - Compset Cal")
            prev_wb.save(prev_path)

            target_wb = openpyxl.Workbook()
            target_wb.create_sheet("Neighborhood Data - Mkt Occ")
            target_wb.create_sheet("Neighborhood Data - Compset Cal")
            copy_pasted_sheets(prev_path, target_wb)

            ws = target_wb["Neighborhood Data - Mkt Occ"]
            self.assertEqual(ws["A1"].value, "Date")
            self.assertEqual(ws["B1"].value, "Market Occupancy")
            self.assertEqual(ws["B2"].value, 42.0)

    def test_noop_when_there_is_no_previous_file(self):
        target_wb = openpyxl.Workbook()
        target_wb.create_sheet("Neighborhood Data - Mkt Occ")
        copy_pasted_sheets(None, target_wb)  # must not raise
        self.assertIsNone(target_wb["Neighborhood Data - Mkt Occ"]["A1"].value)


if __name__ == "__main__":
    unittest.main()
