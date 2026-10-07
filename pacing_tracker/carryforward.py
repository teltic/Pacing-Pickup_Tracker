"""Reads yesterday's generated workbook (if one exists) so a fresh rebuild
never wipes manually-entered Override Request/Notes, and so "New Since Last
Review" can compare today's Signal against what it actually was last time.

The workbook lives in a folder synced by the Google Drive desktop app --
this just reads/writes plain files there, no Drive API/OAuth involved.
"""

import glob
import os
import re
from datetime import date

import openpyxl

FILENAME_RE = re.compile(r"Daily_Pacing_Pickup_(\d{1,2})\.(\d{1,2})\.(\d{2})\.xlsx$")

# 0-based column indices within the Daily Pacing sheet's row tuples.
# 2026-10-07: shifted left (Pickup 3d/14d/30d/60d dropped) -- Signal=L,
# Override Request=M, Notes=N.
COL_DATE, COL_SIGNAL, COL_OVERRIDE_REQUEST, COL_NOTES = 0, 11, 12, 13

# Sheets that hold a raw paste of a PriceLabs dashboard export (see
# excel_report.py's module docstring) rather than anything this tool
# computes -- carried forward verbatim on every regenerate so a fresh
# daily file never wipes out the user's last paste.
PASTED_SHEET_NAMES = ["Neighborhood Data - Mkt Occ", "Neighborhood Data - Compset Cal"]


def parse_filename_date(filename):
    match = FILENAME_RE.search(filename)
    if not match:
        return None
    month, day, year2 = (int(x) for x in match.groups())
    try:
        return date(2000 + year2, month, day)
    except ValueError:
        return None


def _dated_files(folder):
    """[(file_date, path), ...] for every Daily_Pacing_Pickup_*.xlsx in
    folder whose name parses to a date, sorted oldest to newest.
    """
    candidates = []
    for path in glob.glob(os.path.join(folder, "Daily_Pacing_Pickup_*.xlsx")):
        file_date = parse_filename_date(os.path.basename(path))
        if file_date is not None:
            candidates.append((file_date, path))
    candidates.sort()
    return candidates


def find_previous_file(folder, before_date):
    """Most recent Daily_Pacing_Pickup_*.xlsx in folder strictly before
    before_date, or None if there isn't one.
    """
    earlier = [(d, p) for d, p in _dated_files(folder) if d < before_date]
    return earlier[-1][1] if earlier else None


def find_latest_file(folder):
    """Most recent Daily_Pacing_Pickup_*.xlsx in folder (by filename date,
    not filesystem mtime -- Drive sync can touch mtimes), or None if the
    folder has none.
    """
    dated = _dated_files(folder)
    return dated[-1][1] if dated else None


def load_previous_state(path):
    """{date_str: {"override_request", "notes", "signal"}} keyed by ISO
    date, read from a previously-generated (and recalculated -- this needs
    cached formula values, not just formula strings) workbook.
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Daily Pacing"]
    state = {}
    for row in ws.iter_rows(min_row=2):
        date_value = row[COL_DATE].value
        if date_value is None:
            continue
        date_str = date_value.strftime("%Y-%m-%d") if hasattr(date_value, "strftime") else str(date_value)
        state[date_str] = {
            "override_request": row[COL_OVERRIDE_REQUEST].value,
            "notes": row[COL_NOTES].value,
            "signal": row[COL_SIGNAL].value,
        }
    return state


def load_previous_state_for_folder(folder, pull_date):
    """Convenience: find + load in one call. Returns {} if there's no
    previous file (e.g. first run ever), rather than erroring -- a missing
    history is expected on day one, not a failure.
    """
    previous_path = find_previous_file(folder, pull_date)
    if previous_path is None:
        return {}
    return load_previous_state(previous_path)


def copy_pasted_sheets(previous_path, target_wb):
    """Copies every cell of PASTED_SHEET_NAMES from yesterday's workbook
    into a freshly-generated one, verbatim (values, not formulas -- these
    sheets are raw pastes and have none). A no-op if previous_path is None
    (first run ever) or a sheet doesn't exist in the previous file (e.g.
    the very first workbook built under this design). target_wb's sheets
    are expected to already exist (excel_report.build_workbook creates
    them, empty, before calling this).
    """
    if previous_path is None:
        return
    prev_wb = openpyxl.load_workbook(previous_path, data_only=True)
    for sheet_name in PASTED_SHEET_NAMES:
        if sheet_name not in prev_wb.sheetnames:
            continue
        prev_ws = prev_wb[sheet_name]
        target_ws = target_wb[sheet_name] if sheet_name in target_wb.sheetnames else target_wb.create_sheet(sheet_name)
        for row in prev_ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    target_ws.cell(row=cell.row, column=cell.column, value=cell.value)
